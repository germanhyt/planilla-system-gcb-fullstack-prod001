"""Incremento 3 Sprint 3: criterio PT unificado, pantalla Reportes y exportación a Planilla Core."""
import csv
import io
import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.models import Asistencia, Auditoria, BonoNocturno, CierrePeriodo, Empresa, RuleSet, TipoContrato, Trabajador

PER = "2026-04"


def _rs(db, empresa_id):
    return db.query(RuleSet).filter(RuleSet.empresa_id == empresa_id).order_by(RuleSet.valid_from).first().id


def _trab(db, emp_codigo, contrato=None, excluir=()):
    q = db.query(Trabajador).join(Empresa, Empresa.id == Trabajador.empresa_id).filter(Empresa.codigo == emp_codigo,
                                                                                    Trabajador.codigo.notin_(excluir))
    if contrato:
        q = q.join(TipoContrato, TipoContrato.id == Trabajador.tipo_contrato_id).filter(TipoContrato.codigo == contrato)
    return q.order_by(Trabajador.codigo.desc()).first()


@pytest.fixture(scope="module")
def datos():
    """Asistencias ficticias de abril 2026 para un trabajador CAVAS y uno GCB SA (incluye cruce de mes)."""
    from app.db import SessionLocal
    db = SessionLocal()
    cav = _trab(db, "CAVAS", "PT85", excluir=("CAV-007",))
    otra = db.query(Empresa).filter(Empresa.codigo != "CAVAS").order_by(Empresa.id).first()
    gcb = _trab(db, otra.codigo)
    for t in (cav, gcb):
        rs = _rs(db, t.empresa_id)
        # 3 días x 5 h la semana del 6-abr (lunes): promedio 5 h/día trabajado (con /7 serían 2.14)
        for i in range(3):
            db.add(Asistencia(trabajador_id=t.id, dia_operativo=date(2026, 4, 6) + timedelta(days=i), periodo=PER,
                              horas_programadas=4, horas_trabajadas=5, tardanza_min=10, horas_extra_25=1,
                              horas_extra_35=0.5, horas_nocturnas=2, rule_set_id=rs))
        # turno que inicia el 30-abr 22:00 y termina el 1-may: pertenece a abril por día operativo
        db.add(Asistencia(trabajador_id=t.id, dia_operativo=date(2026, 4, 30), periodo=PER, horas_programadas=8,
                          horas_trabajadas=8, horas_nocturnas=8, rule_set_id=rs))
        # 1-may no entra en abril
        db.add(Asistencia(trabajador_id=t.id, dia_operativo=date(2026, 5, 1), periodo="2026-05", horas_programadas=8,
                          horas_trabajadas=8, rule_set_id=rs))
    adm = db.execute(__import__("sqlalchemy").text("select id from usuario where email='admin@demo.gcb.pe'")).scalar()
    db.add(BonoNocturno(trabajador_id=cav.id, periodo=PER, cantidad=2, estado="aprobado", registrado_por=adm))
    db.commit()
    out = {"cav": cav.codigo, "gcb": gcb.codigo, "cav_id": cav.id, "cav_emp": cav.empresa_id, "gcb_emp": gcb.empresa_id,
           "gcb_unidad": gcb.unidad_id, "cav_unidad": cav.unidad_id}
    db.close()
    return out


def test_pt_riesgo_promedio_por_dia_trabajado(client, H, datos):
    r = client.get(f"/excepciones/part-time-riesgo?periodo={PER}", headers=H["rrhh"]).json()
    a = next(x for x in r if x["codigo"] == datos["cav"] and x["semana_inicio"] == "2026-04-06")
    assert a["dias_trabajados"] == 3 and a["total_horas_semana"] == 15
    assert a["promedio_diario"] == 5.0 and a["bloqueante"] is False
    # La bandeja usa el mismo criterio
    b = client.get("/excepciones/bandeja", params={"desde": "2026-04-06", "hasta": "2026-04-12", "tipo": "pt_alerta"},
                   headers=H["rrhh"]).json()
    assert any(i["trabajador"]["codigo"] == datos["cav"] and "5.0" in i["detalle"] for i in b)


def test_consolidado_por_dia_operativo_y_recargo_solo_cavas(client, H, datos):
    r = client.get("/reportes/consolidado-mensual", params={"periodo": PER, "empresa_id": datos["cav_emp"]},
                   headers=H["rrhh"]).json()
    assert r["muestra_recargo_nocturno"] is True and r["periodo_cerrado"] is False
    c = next(x for x in r["registros"] if x["codigo"] == datos["cav"])
    assert c["horas_trabajadas"] == 23.0  # 3x5 + 8 (30-abr); el 1-may no entra
    assert c["horas_extra_25"] == 3.0 and c["horas_extra_35"] == 1.5 and c["horas_normales"] == 18.5
    assert c["tardanza_min"] == 30 and c["horas_nocturnas"] == 14.0 and c["recargo_nocturno_horas"] == 14.0
    assert c["bono_3am"] == 2
    g = client.get("/reportes/consolidado-mensual", params={"periodo": PER, "empresa_id": datos["gcb_emp"]},
                   headers=H["rrhh"]).json()
    assert g["muestra_recargo_nocturno"] is False
    x = next(x for x in g["registros"] if x["codigo"] == datos["gcb"])
    assert x["horas_nocturnas"] == 14.0 and x["recargo_nocturno_horas"] is None


def test_consolidado_jefe_solo_su_unidad(client, H, datos):
    r = client.get("/reportes/consolidado-mensual", params={"periodo": PER}, headers=H["jefe"])
    assert r.status_code == 200
    unidades = {x["unidad_id"] for x in r.json()["registros"]}
    assert unidades <= {datos["cav_unidad"]}
    if datos["gcb_unidad"] != datos["cav_unidad"]:
        assert client.get("/reportes/consolidado-mensual", params={"periodo": PER, "unidad_id": datos["gcb_unidad"]},
                          headers=H["jefe"]).status_code == 403
    assert client.post("/reportes/export-planilla-core", params={"periodo": PER, "empresa_id": datos["cav_emp"]},
                       headers=H["jefe"]).status_code == 403


def test_descargas_csv_xlsx(client, H, datos):
    p = {"periodo": PER, "empresa_id": datos["cav_emp"]}
    r = client.get("/reportes/consolidado-mensual", params=p | {"formato": "csv"}, headers=H["rrhh"])
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    filas = list(csv.reader(io.StringIO(r.text.lstrip("\ufeff"))))
    assert filas[0][0] == "Código" and any(f[0] == datos["cav"] for f in filas[1:])
    r = client.get("/reportes/consolidado-mensual", params=p | {"formato": "xlsx"}, headers=H["rrhh"])
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(1, 1).value == "Código" and ws.max_row >= 2


def test_export_preliminar_luego_final_versionado_y_auditado(client, H, datos, db):
    p = {"periodo": PER, "empresa_id": datos["cav_emp"]}
    assert client.post("/reportes/export-planilla-core", params=p | {"solo_cerrado": True},
                       headers=H["rrhh"]).status_code == 409
    r = client.post("/reportes/export-planilla-core", params=p, headers=H["rrhh"])
    assert r.status_code == 200 and r.headers["x-export-estado"] == "PRELIMINAR"
    assert "PRELIMINAR" in r.headers["content-disposition"] and "_v1_" in r.headers["content-disposition"]
    filas = list(csv.DictReader(io.StringIO(r.text), delimiter=";"))
    mias = {f["concepto_codigo"]: f for f in filas if f["trabajador_codigo"] == datos["cav"]}
    assert mias["HE25"]["cantidad"] == "3.0" and mias["RNOC"]["cantidad"] == "14.0" and mias["B3AM"]["cantidad"] == "2"
    assert all(f["layout_version"] == "PC-ASIS-1.0" and f["estado"] == "PRELIMINAR" for f in filas)
    # cerrar periodo -> FINAL v2 (JSON)
    db.add(CierrePeriodo(empresa_id=datos["cav_emp"], periodo=PER, cerrado_por=1, estado="cerrado"))
    db.commit()
    r = client.post("/reportes/export-planilla-core", params=p | {"formato": "json", "solo_cerrado": True},
                    headers=H["rrhh"])
    assert r.status_code == 200
    j = json.loads(r.content)
    assert j["estado"] == "FINAL" and j["version_archivo"] == 2 and "RNOC" in j["conceptos"]
    aud = db.query(Auditoria).filter(Auditoria.entidad == "export_planilla_core").all()
    assert len([a for a in aud if a.entidad_id.endswith(PER)]) >= 2


def test_export_gcb_sin_recargo_nocturno(client, H, datos):
    r = client.post("/reportes/export-planilla-core", params={"periodo": PER, "empresa_id": datos["gcb_emp"]},
                    headers=H["admin"])
    filas = list(csv.DictReader(io.StringIO(r.text), delimiter=";"))
    cods = {f["concepto_codigo"] for f in filas if f["trabajador_codigo"] == datos["gcb"]}
    assert "HNOC" in cods and "RNOC" not in cods